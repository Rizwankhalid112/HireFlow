"""Apply accepted keyword rewrites to a CV.

The narrow operation this module exists to do: swap *the specific words the user
approved*, and change nothing else. Not a rewrite, not a regeneration — the
output has to be recognisably the same document, in the same template, with the
same layout, differing only in the words on the change list.

Three rules hold everywhere below, because a wrong edit on someone's CV is worse
than no edit at all:

1. **Verify before replacing.** If the CV does not contain the wording verbatim,
   the rewrite is skipped and reported. Nothing is fuzzy-matched.
2. **Word boundaries only.** Replacing `Postgres` must never turn an existing
   `PostgreSQL` into `PostgreSQLSQL`.
3. **Nothing is silent.** Every rewrite ends up in `applied` or in `skipped`
   with a reason the user can read.

Only the `reworded` bucket is ever applied. `missing` is a question — applying it
would mean writing a claim onto a document the user has to defend in an
interview.
"""

import io
import logging
import re

logger = logging.getLogger(__name__)

# Guard against a pathological document turning one rewrite into a long loop.
MAX_OCCURRENCES_PER_PARAGRAPH = 50


def _matcher(wording):
    """Word-boundary match that also survives punctuation in the term.

    `\\b` is wrong here: it is defined against word characters, so it fails on
    the terms this feature most needs — `C++`, `.NET`, `Node.js`. Lookarounds
    for a word character on either side handle all of them.
    """
    return re.compile(rf'(?<!\w){re.escape(wording)}(?!\w)', re.IGNORECASE)


def plan(cv_text, rewrites):
    """What would change, before anything does. `(planned, skipped)`.

    Called before the user is asked to accept, so the confirmation screen shows
    the true number of occurrences rather than a guess, and so a rewrite that
    cannot be applied is visible up front instead of quietly missing afterwards.
    """
    planned, skipped = [], []

    for rewrite in rewrites or []:
        keyword = (rewrite.get('keyword') or '').strip()
        current = (rewrite.get('current_wording') or '').strip()
        record = {
            'keyword': keyword,
            'current_wording': current,
            'where': (rewrite.get('where') or '').strip(),
        }

        if not keyword or not current:
            skipped.append({**record, 'reason': 'Incomplete suggestion.'})
            continue

        if keyword.casefold() == current.casefold():
            skipped.append({**record, 'reason': 'The CV already uses this wording.'})
            continue

        found = len(_matcher(current).findall(cv_text or ''))
        if not found:
            skipped.append({
                **record,
                'reason': f'“{current}” does not appear in this CV.',
            })
            continue

        planned.append({**record, 'occurrences': found})

    return _drop_overlapping(planned, skipped)


def _drop_overlapping(planned, skipped):
    """Refuse two rewrites that touch the same words.

    Seen in testing: alongside `Postgres → PostgreSQL`, the model proposed
    `Optimised Postgres queries → query optimisation`. That second one is not a
    rename, it is prose being rewritten, and applying both would have produced
    “query optimisation, cutting a report endpoint from 4.2s to 310ms.”

    It happened to be skipped only because the first rewrite changed the text
    out from under it — order deciding whether a CV is mangled is not a
    guarantee, so overlap is now refused outright.

    The shorter wording wins. A rewrite that spans a phrase is the prose case;
    the term-level rename is the one this feature exists to make.
    """
    kept = []

    for item in sorted(planned, key=lambda r: len(r['current_wording'])):
        current = item['current_wording'].casefold()
        clash = next(
            (k for k in kept
             if current in k['current_wording'].casefold()
             or k['current_wording'].casefold() in current),
            None,
        )
        if clash:
            skipped.append({
                **item,
                'reason': f'Overlaps “{clash["current_wording"]}”, which we are '
                          'already changing.',
            })
        else:
            kept.append(item)

    return kept, skipped


# --- Plain text ---------------------------------------------------------------

def apply_to_text(text, planned):
    """The reference implementation, and what the change list is computed from."""
    for rewrite in planned:
        text = _matcher(rewrite['current_wording']).sub(rewrite['keyword'], text)
    return text


# --- DOCX: edit the user's own file in place ----------------------------------

def _iter_paragraphs(document):
    """Body, tables, headers and footers.

    Tables are recursive — a table cell can hold another table — and a CV built
    from a two-column Word template keeps *everything* in table cells, so
    skipping them would silently do nothing on exactly the documents people use.
    """
    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from walk_table(nested)

    yield from document.paragraphs
    for table in document.tables:
        yield from walk_table(table)

    for section in document.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is None:
                continue
            yield from part.paragraphs
            for table in part.tables:
                yield from walk_table(table)


def _replace_in_paragraph(paragraph, pattern, replacement):
    """Replace every occurrence in one paragraph. Returns how many.

    Word splits a word across runs at any formatting or spell-check boundary —
    `Amazon Web Services` is commonly three runs — so a run-by-run search misses
    the cases that matter. The match is therefore made against the paragraph's
    concatenated text, and written back into the first run it overlaps, with the
    matched span cleared from the rest.

    Formatting survives because it lives on the run, and the runs themselves are
    never replaced: only their text changes.
    """
    replaced = 0

    for _ in range(MAX_OCCURRENCES_PER_PARAGRAPH):
        runs = paragraph.runs
        if not runs:
            return replaced

        match = pattern.search(''.join(run.text for run in runs))
        if match is None:
            return replaced

        start, end = match.span()
        position, written = 0, False

        for run in runs:
            run_start, run_end = position, position + len(run.text)
            position = run_end

            # Does this run overlap the matched span at all?
            if run_end <= start or run_start >= end:
                continue

            head = run.text[:max(0, start - run_start)]
            tail = run.text[max(0, end - run_start):] if run_end > end else ''
            run.text = head + ('' if written else replacement) + tail
            written = True

        replaced += 1

    logger.warning('Stopped replacing after %d occurrences in one paragraph',
                   MAX_OCCURRENCES_PER_PARAGRAPH)
    return replaced


def apply_to_docx(blob, planned):
    """Edit the user's own DOCX. `(bytes, applied, skipped)`.

    His template, his fonts, his layout — only the approved words differ.
    """
    from docx import Document

    document = Document(io.BytesIO(blob))
    applied, skipped = [], []

    for rewrite in planned:
        pattern = _matcher(rewrite['current_wording'])
        count = sum(
            _replace_in_paragraph(paragraph, pattern, rewrite['keyword'])
            for paragraph in _iter_paragraphs(document)
        )

        if count:
            applied.append({**rewrite, 'occurrences': count})
        else:
            # Planned against the extracted text but not found in the document
            # itself — the extractor flattens things Word keeps apart, such as a
            # term split over a line break.
            skipped.append({
                **rewrite,
                'reason': f'“{rewrite["current_wording"]}” could not be located '
                          'in the document layout.',
            })

    out = io.BytesIO()
    document.save(out)
    return out.getvalue(), applied, skipped


# --- HireFlow-built CV: re-render through the user's own template -------------

# `template` is the registry entry — file paths and flags, not user content.
_CONTEXT_SKIP_KEYS = frozenset({'template'})


def _rewrite_value(value, pairs):
    if isinstance(value, str):
        for pattern, keyword in pairs:
            value = pattern.sub(keyword, value)
        return value
    if isinstance(value, list):
        return [_rewrite_value(item, pairs) for item in value]
    if isinstance(value, dict):
        return {key: _rewrite_value(item, pairs) for key, item in value.items()}
    return value


def apply_to_profile(cv, planned, template_id=None):
    """`(pdf_bytes, content_json, applied, skipped, resolved_template_id)`.

    Renders through the template the user already chose. Nothing is written to
    the database — the master CV is not the thing being tailored, and a tailored
    CV is a snapshot beside it.
    """
    from apps.cv_builder.services.cv_context import build_cv_context
    from apps.cv_builder.services.pdf_renderer import render_context_pdf
    from apps.cv_builder.templates_registry import resolve_template, resolve_template_id

    requested = template_id or cv.template_id
    template = resolve_template(requested)
    resolved_id = resolve_template_id(requested)

    context = build_cv_context(cv, template)
    pairs = [(_matcher(r['current_wording']), r['keyword']) for r in planned]

    tailored = {
        key: (value if key in _CONTEXT_SKIP_KEYS else _rewrite_value(value, pairs))
        for key, value in context.items()
    }

    # Counted from the rendered content rather than trusted from the plan: the
    # plan was made against the flattened CV text, and the context is the thing
    # actually being printed.
    before = _content_text(context)
    applied, skipped = [], []
    for rewrite in planned:
        found = len(_matcher(rewrite['current_wording']).findall(before))
        if found:
            applied.append({**rewrite, 'occurrences': found})
        else:
            skipped.append({
                **rewrite,
                'reason': f'“{rewrite["current_wording"]}” is not in the printed CV.',
            })

    content_json = {k: v for k, v in tailored.items() if k not in _CONTEXT_SKIP_KEYS}
    return render_context_pdf(tailored, template), content_json, applied, skipped, resolved_id


def _content_text(context):
    """Every string in the context, flattened — for counting only."""
    parts = []

    def walk(value):
        if isinstance(value, str):
            parts.append(value)
        elif isinstance(value, list):
            for item in value:
                walk(item)
        elif isinstance(value, dict):
            for key, item in value.items():
                if key not in _CONTEXT_SKIP_KEYS:
                    walk(item)

    walk({k: v for k, v in context.items() if k not in _CONTEXT_SKIP_KEYS})
    return '\n'.join(parts)
