"""Stage 1: file to text.

Fixtures are generated at test time rather than committed as binaries —
WeasyPrint is already a dependency and can produce a PDF, python-docx a DOCX —
so there are no opaque blobs in the repo and the inputs are readable in the test
that uses them.

No database and no network: this module is a pure function of a file on disk.
"""

import pytest
from django.conf import settings

from apps.cv_builder.services.extraction import (
    ExtractionError,
    extract_docx,
    extract_pdf,
    extract_text,
)


def write_pdf(path, html):
    from weasyprint import HTML

    HTML(string=html).write_pdf(str(path))
    return path


def write_docx(path, paragraphs=(), table_rows=()):
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(str(path))
    return path


class TestPdf:
    def test_text_comes_back(self, tmp_path):
        write_pdf(tmp_path / 'cv.pdf', '<h1>Ada Lovelace</h1><p>Backend Engineer</p>')
        text = extract_pdf(str(tmp_path / 'cv.pdf'))
        assert 'Ada Lovelace' in text
        assert 'Backend Engineer' in text

    def test_a_pdf_with_no_text_comes_back_empty(self, tmp_path):
        """How a scanned CV looks. Empty is not an error here — the caller
        decides what it means, and stops before spending an API call."""
        write_pdf(tmp_path / 'blank.pdf', '<div style="height:200px"></div>')
        assert extract_pdf(str(tmp_path / 'blank.pdf')).strip() == ''

    def test_a_corrupt_file_raises_something_actionable(self, tmp_path):
        path = tmp_path / 'broken.pdf'
        path.write_bytes(b'%PDF-1.4 and then nothing that parses')
        with pytest.raises(ExtractionError) as caught:
            extract_pdf(str(path))
        assert 're-export' in str(caught.value).lower()

    def test_a_missing_file_raises_rather_than_crashing_the_worker(self, tmp_path):
        with pytest.raises(ExtractionError):
            extract_pdf(str(tmp_path / 'not-here.pdf'))

    def test_the_page_cap_holds(self, tmp_path, settings):
        settings.CV_EXTRACT_MAX_PAGES = 1
        html = '<p>PAGE ONE</p><div style="page-break-before:always"></div><p>PAGE TWO</p>'
        write_pdf(tmp_path / 'long.pdf', html)
        text = extract_pdf(str(tmp_path / 'long.pdf'))
        assert 'PAGE ONE' in text
        assert 'PAGE TWO' not in text

    def test_the_character_cap_holds(self, tmp_path, settings):
        settings.CV_EXTRACT_MAX_CHARS = 20
        write_pdf(tmp_path / 'cv.pdf', '<p>' + ('word ' * 200) + '</p>')
        assert len(extract_pdf(str(tmp_path / 'cv.pdf'))) == 20


class TestDocx:
    def test_paragraphs_come_back(self, tmp_path):
        write_docx(tmp_path / 'cv.docx', paragraphs=['Ada Lovelace', 'Backend Engineer'])
        text = extract_docx(str(tmp_path / 'cv.docx'))
        assert 'Ada Lovelace' in text
        assert 'Backend Engineer' in text

    def test_table_cells_come_back(self, tmp_path):
        """Two-column CV layouts are usually one table, which puts the entire
        left-hand column outside `document.paragraphs`. Missing this loses the
        contact details and skills of every CV built that way."""
        write_docx(
            tmp_path / 'cv.docx',
            paragraphs=['Profile'],
            table_rows=[['SKILLS', 'EXPERIENCE'], ['Python, Django', 'Acme Corp']],
        )
        text = extract_docx(str(tmp_path / 'cv.docx'))
        assert 'Python, Django' in text
        assert 'Acme Corp' in text

    def test_paragraphs_and_cells_do_not_double_count(self, tmp_path):
        write_docx(tmp_path / 'cv.docx', paragraphs=['Unique Line'], table_rows=[['Cell']])
        assert extract_docx(str(tmp_path / 'cv.docx')).count('Unique Line') == 1

    def test_an_image_only_docx_comes_back_empty(self, tmp_path):
        write_docx(tmp_path / 'empty.docx')
        assert extract_docx(str(tmp_path / 'empty.docx')).strip() == ''

    def test_a_non_docx_raises_something_actionable(self, tmp_path):
        path = tmp_path / 'fake.docx'
        path.write_bytes(b'PK\x03\x04 not really a word file')
        with pytest.raises(ExtractionError) as caught:
            extract_docx(str(path))
        assert 'docx' in str(caught.value).lower()


class TestDispatch:
    def test_it_routes_on_the_stored_type(self, tmp_path):
        write_docx(tmp_path / 'cv.docx', paragraphs=['Routed correctly'])
        assert 'Routed correctly' in extract_text(str(tmp_path / 'cv.docx'), 'docx')

    def test_an_unknown_type_raises(self, tmp_path):
        with pytest.raises(ExtractionError):
            extract_text(str(tmp_path / 'whatever'), 'rtf')
